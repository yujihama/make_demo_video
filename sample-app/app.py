"""撮影検証用の対象アプリ: 証跡をアップロード → 監査実行 → 結果表示 → Excel ダウンロード。"""
import asyncio
import csv
import io
import time
import uuid

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

app = FastAPI()
JOBS: dict[str, dict] = {}

RULES = [
    ("承認者未記載", lambda r: not r.get("承認者")),
    ("自己承認", lambda r: r.get("承認者") and r.get("承認者") == r.get("申請者")),
    ("高額取引（100万円超）", lambda r: _amount(r) > 1_000_000),
]


def _amount(row: dict) -> int:
    try:
        return int(str(row.get("金額", "0")).replace(",", ""))
    except ValueError:
        return 0


def _read_rows(name: str, data: bytes) -> list[dict]:
    if name.lower().endswith(".xlsx"):
        ws = load_workbook(io.BytesIO(data), read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        return [{h: ("" if v is None else str(v)) for h, v in zip(header, r)} for r in rows[1:]]
    text = data.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


@app.post("/api/jobs")
async def create_job(file: UploadFile = File(...)):
    rows = _read_rows(file.filename, await file.read())
    if not rows:
        raise HTTPException(400, "証跡ファイルに明細がありません")
    job_id = uuid.uuid4().hex[:8]
    JOBS[job_id] = {"file": file.filename, "rows": rows, "started": time.time()}
    return {"job_id": job_id, "rows": len(rows)}


@app.get("/api/jobs/{job_id}")
def job_status(job_id: str):
    job = JOBS.get(job_id) or {}
    if not job:
        raise HTTPException(404)
    # 監査エージェントの処理時間を模す（約4秒）
    progress = min(100, int((time.time() - job["started"]) / 4 * 100))
    if progress < 100:
        return {"status": "running", "progress": progress}
    findings = []
    for i, row in enumerate(job["rows"], start=1):
        for name, rule in RULES:
            if rule(row):
                findings.append({"no": i, "id": row.get("伝票番号", ""), "rule": name,
                                 "amount": _amount(row), "applicant": row.get("申請者", "")})
    job["findings"] = findings
    return {"status": "done", "progress": 100, "rows": len(job["rows"]), "findings": findings}


@app.get("/api/jobs/{job_id}/report.xlsx")
def report(job_id: str):
    job = JOBS.get(job_id)
    if not job or "findings" not in job:
        raise HTTPException(404)
    wb = Workbook()
    ws = wb.active
    ws.title = "監査結果"
    ws.append(["No", "伝票番号", "指摘内容", "金額", "申請者"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for f in job["findings"]:
        ws.append([f["no"], f["id"], f["rule"], f["amount"], f["applicant"]])
    for col, w in zip("ABCDE", (6, 14, 26, 14, 12)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="audit_report_{job_id}.xlsx"'},
    )


@app.get("/health")
def health():
    return {"ok": True}


# --- 2つ目のモジュール: 規程Q&A（P6 の「YAML 追加だけで撮れるか」の検証用） ---
ANSWERS = {
    "接待": "接待交際費は1人あたり1万円を超える場合、事前に部長承認が必要です（経費規程 第12条）。",
    "承認": "申請者本人による承認（自己承認）は認められません。上長または代理承認者が承認します（経費規程 第8条）。",
}


@app.post("/api/ask")
async def ask(q: dict):
    await asyncio.sleep(1.5)  # 回答生成を模す
    text = q.get("question", "")
    answer = next((a for k, a in ANSWERS.items() if k in text), "該当する規程が見つかりませんでした。")
    return {"answer": answer}


@app.get("/policy", response_class=HTMLResponse)
def policy():
    return POLICY_PAGE


POLICY_PAGE = """<!doctype html>
<html lang="ja"><head><meta charset="utf-8"><title>規程Q&A - 内部監査エージェント</title>
<style>
 body{font-family:"Noto Sans CJK JP","Yu Gothic UI","Meiryo",sans-serif;margin:0;background:#f4f6f9;color:#1d2533}
 header{background:#1f4e78;color:#fff;padding:14px 32px;font-size:20px;font-weight:600}
 main{max-width:1100px;margin:28px auto;padding:0 24px}
 .card{background:#fff;border-radius:10px;box-shadow:0 1px 4px #0002;padding:24px;margin-bottom:20px}
 h2{margin:0 0 14px;font-size:18px}
 textarea{width:100%;height:90px;font-size:16px;padding:10px;border:1px solid #9fb3c8;border-radius:6px;box-sizing:border-box;font-family:inherit}
 button{background:#1f4e78;color:#fff;border:0;border-radius:6px;padding:10px 22px;font-size:15px;cursor:pointer;margin-top:12px}
 .answer{border-left:4px solid #2f80ed;background:#f0f6ff;padding:14px 18px;font-size:16px;line-height:1.7}
 .hidden{display:none}
</style></head>
<body>
<header data-testid="app-title">内部監査エージェント（デモ） / 規程Q&A</header>
<main>
 <section class="card">
  <h2>社内規程について質問する</h2>
  <textarea data-testid="question-input" id="q" placeholder="例: 接待交際費の上限は？"></textarea>
  <button data-testid="ask-button" id="ask">質問する</button>
  <div data-testid="ask-status" id="st"></div>
 </section>
 <section class="card hidden" id="res" data-testid="answer-card">
  <h2>回答</h2>
  <div class="answer" data-testid="answer" id="ans"></div>
 </section>
</main>
<script>
const $=id=>document.getElementById(id);
$('ask').addEventListener('click',async()=>{
  $('st').textContent='規程を検索しています…';
  const r=await (await fetch('/api/ask',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({question:$('q').value})})).json();
  $('st').textContent='回答しました';$('ans').textContent=r.answer;$('res').classList.remove('hidden');
});
</script>
</body></html>"""


@app.get("/", response_class=HTMLResponse)
def index():
    return PAGE


PAGE = """<!doctype html>
<html lang="ja"><head><meta charset="utf-8"><title>内部監査エージェント</title>
<style>
 body{font-family:"Noto Sans CJK JP","Yu Gothic UI","Meiryo",sans-serif;margin:0;background:#f4f6f9;color:#1d2533}
 header{background:#1f4e78;color:#fff;padding:14px 32px;font-size:20px;font-weight:600}
 main{max-width:1100px;margin:28px auto;padding:0 24px}
 .card{background:#fff;border-radius:10px;box-shadow:0 1px 4px #0002;padding:24px;margin-bottom:20px}
 h2{margin:0 0 14px;font-size:18px}
 .drop{border:2px dashed #9fb3c8;border-radius:8px;padding:22px;text-align:center;color:#52606d}
 button,.btn{background:#1f4e78;color:#fff;border:0;border-radius:6px;padding:10px 22px;font-size:15px;cursor:pointer;text-decoration:none;display:inline-block}
 button:disabled{background:#9aa5b1;cursor:default}
 .bar{height:12px;background:#e4e7eb;border-radius:6px;overflow:hidden;margin:12px 0}
 .bar>div{height:100%;background:#2f80ed;width:0;transition:width .3s}
 table{border-collapse:collapse;width:100%;font-size:14px}
 th,td{border-bottom:1px solid #e4e7eb;padding:8px 10px;text-align:left}
 th{background:#f0f4f8}
 .tag{background:#fde8e8;color:#b42318;border-radius:4px;padding:2px 8px;font-size:13px}
 .summary{font-size:15px;margin-bottom:12px}
 .hidden{display:none}
</style></head>
<body>
<header data-testid="app-title">内部監査エージェント（デモ）</header>
<main>
 <section class="card" data-testid="upload-card">
  <h2>1. 証跡ファイルのアップロード</h2>
  <div class="drop">
   <p>経費精算の明細（CSV / Excel）を選択してください</p>
   <input type="file" accept=".csv,.xlsx" data-testid="evidence-input" id="evidence">
   <p data-testid="file-name" id="fileName"></p>
  </div>
 </section>
 <section class="card">
  <h2>2. 監査の実行</h2>
  <button data-testid="run-button" id="run" disabled>監査を実行</button>
  <div class="bar"><div id="bar" data-testid="progress-bar"></div></div>
  <div data-testid="status" id="status">待機中</div>
 </section>
 <section class="card hidden" id="result" data-testid="result-card">
  <h2>3. 監査結果</h2>
  <div class="summary" data-testid="summary" id="summary"></div>
  <table data-testid="findings-table"><thead><tr><th>No</th><th>伝票番号</th><th>指摘内容</th><th>金額</th><th>申請者</th></tr></thead><tbody id="rows"></tbody></table>
  <p><a class="btn" data-testid="download-button" id="download" href="#">結果をExcelでダウンロード</a></p>
 </section>
</main>
<script>
const $=id=>document.getElementById(id);
$('evidence').addEventListener('change',e=>{const f=e.target.files[0];$('fileName').textContent=f?`選択中: ${f.name}`:'';$('run').disabled=!f;});
$('run').addEventListener('click',async()=>{
  $('run').disabled=true;$('status').textContent='アップロード中…';
  const fd=new FormData();fd.append('file',$('evidence').files[0]);
  const r=await (await fetch('/api/jobs',{method:'POST',body:fd})).json();
  $('status').textContent=`${r.rows}件の明細を監査しています…`;
  const poll=async()=>{
    const s=await (await fetch(`/api/jobs/${r.job_id}`)).json();
    $('bar').style.width=s.progress+'%';
    if(s.status!=='done'){setTimeout(poll,400);return;}
    $('status').textContent='監査が完了しました';$('status').dataset.done='true';
    $('summary').textContent=`${s.rows}件中 ${s.findings.length}件の指摘があります`;
    $('rows').innerHTML=s.findings.map(f=>`<tr><td>${f.no}</td><td>${f.id}</td><td><span class="tag">${f.rule}</span></td><td>${f.amount.toLocaleString()}円</td><td>${f.applicant}</td></tr>`).join('');
    $('download').href=`/api/jobs/${r.job_id}/report.xlsx`;
    $('result').classList.remove('hidden');
  };poll();
});
</script>
</body></html>"""
